"""The mixed-format library for retrieval set M01–M12 (eval/retrieval_set_20261024.yaml).

Word (via macOS textutil), PowerPoint (python-pptx, with speaker notes), Excel
(openpyxl, two sheets), a saved web page whose script holds a word the text does
not, Markdown notes, a scanned PDF and a sideways phone photo (both read by local
OCR), and a text PDF whose answer is on page 5; plus distractors on nearby topics.
Needs macOS (textutil, Arial Unicode, Vision). Frozen with the set: changing a file
changes the answers.
"""
import subprocess
import tempfile
from pathlib import Path

from eval.ocr_cases import _picture, _scanned_pdf


def _docx(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        source = Path(tmp) / 'note.txt'
        source.write_text(text, encoding='utf-8')
        subprocess.run(['textutil', '-convert', 'docx', '-inputencoding', 'UTF-8', str(source), '-output', str(path)],
                       check=True, capture_output=True)


def _pptx(path, slides):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    for title, body, notes in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4)).text_frame.text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)


def _xlsx(path, sheets):
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _pdf(path, pages):
    import pymupdf
    path.parent.mkdir(parents=True, exist_ok=True)
    with pymupdf.open() as document:
        for text in pages:
            document.new_page().insert_textbox(pymupdf.Rect(40, 40, 550, 800), text, fontsize=11, fontname='china-s')
        document.save(path)


def _write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding='utf-8')


PAGE = ('<html><head><style>.ad{color:red}</style>'
        '<script>var trackingId = "quantumleap"; function load(){}</script></head><body>'
        '<h1>知识蒸馏入门</h1><p>知识蒸馏是一种模型压缩方法：学生模型模仿教师模型输出的软标签。</p>'
        '<p>蒸馏时用温度参数 T 软化概率分布，T 越大分布越平滑。</p></body></html>')


def build(root):
    root = Path(root)
    _docx(root / '会议/2026-09-组会纪要.docx',
          '组会纪要 2026 年 9 月 18 日\n参会：王老师、李明、陈静\n'
          '1. 李明汇报了谣言传播图的对比学习方案，图增强采用随机删边。\n'
          '2. 下次组会由陈静汇报消融实验结果。\n3. 服务器预约：每周三晚上统一排队。\n')
    _docx(root / '会议/2026-10-组会纪要.docx',
          '组会纪要 2026 年 10 月 9 日\n参会：王老师、李明、陈静\n'
          '1. 讨论 ACL 投稿截止时间。\n2. 李明负责打印海报。\n')
    _pptx(root / '汇报/期中汇报.pptx', [
        ('期中汇报：多模态谣言检测', '汇报人：陈静', ''),
        ('研究背景', '社交媒体上的谣言常常配有图片', ''),
        ('实验设置', '数据集：Weibo21\n学习率 2e-5，batch size 32，训练 10 轮', ''),
        ('下一步计划', '补充跨领域实验', '备注：答辩时说明为什么没有用 CLIP 特征'),
    ])
    _pptx(root / '汇报/开题报告.pptx', [
        ('开题报告：联邦学习在医疗数据上的应用', '汇报人：李明', ''),
        ('研究计划', '第一年完成医院数据接入', '备注：经费预算另附'),
    ])
    _xlsx(root / '数据/实验结果.xlsx', [
        ('主实验', [('模型', 'Accuracy', 'F1'), ('BERT', 0.861, 0.852), ('RoBERTa', 0.874, 0.866), ('本文方法', 0.902, 0.897)]),
        ('超参数', [('参数', '取值'), ('温度系数', 0.07), ('学习率', '2e-5'), ('批大小', 32)]),
    ])
    _xlsx(root / '数据/经费.xlsx', [('2026', [('项目', '金额'), ('设备费', 12000), ('打印费', 300)])])
    _write(root / '网页/知识蒸馏入门.html', PAGE)
    _write(root / '笔记/读书笔记.md',
           '# 读书笔记\n\n图神经网络层数加深后会出现过平滑：不同节点的表示趋于相同。\n'
           '缓解办法：残差连接、DropEdge。\n')
    _write(root / '笔记/待办.md', '# 待办\n\n- 预约服务器\n- 整理参考文献\n')
    _scanned_pdf(root / '扫描/报销单.pdf', ['差旅费报销单', '出差地点：杭州', '高铁票 金额 553 元'])
    _picture(root / '照片/白板.jpg', ['下周计划', '整理 FakeNewsNet 数据', '复现 GossipCop 基线'], turn=True)
    _pdf(root / '报告/年度报告.pdf', [
        '实验室年度报告 2026\n本报告总结实验室一年来的研究进展。',
        '一、研究方向\n谣言检测、多模态学习与知识图谱。',
        '二、论文发表\n全年发表会议论文六篇、期刊论文两篇。',
        '三、人才培养\n毕业博士一名、硕士三名。',
        '附录 B：实验服务器配置\n计算节点配有 4 张 A100 显卡，每张显存 80GB，共享存储 200TB。',
        '附录 C：致谢\n感谢学院对实验室建设的支持。',
    ])
    return root
