"""Reading Office documents, web pages and reStructuredText (plan 2026-10-24, stage 1)."""
import json
import shutil
import subprocess

import pytest

from agent import extract
from agent.artifacts import ArtifactStore
from agent.task_runtime import TaskRuntime
from safety.sandbox import Sandbox
from tools import fs_read, registry  # noqa: F401  (fs_read registers read_file)

needs_textutil = pytest.mark.skipif(not shutil.which('textutil'), reason='macOS textutil')


def library(tmp_path, **cfg):
    sb = Sandbox(tmp_path); sb.artifacts = ArtifactStore(sb)
    sb.runtime = TaskRuntime('读取', dict(protocol='v3', **cfg), sb, 'one')
    return sb


def read(sb, path, **kw):
    return registry.dispatch('read_file', json.dumps(dict(path=path, **kw)), sb)


def word(path, text, fmt='docx'):
    source = path.with_suffix('.txt')
    source.write_text(text)
    subprocess.run(['textutil', '-convert', fmt, '-output', str(path), str(source)], check=True)
    source.unlink()


def slides(path):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = 'Agent 架构图'
    first.placeholders[1].text = '规划、执行、核验三层'
    first.notes_slide.notes_text_frame.text = '讲到核验时停顿'
    second = deck.slides.add_slide(deck.slide_layouts[5])
    table = second.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    for (r, c), v in {(0, 0): '指标', (0, 1): '数值', (1, 0): 'Recall@5', (1, 1): '81.8%'}.items():
        table.cell(r, c).text = v
    group = second.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text_frame.text = '分组里的说明'
    deck.save(path)


def sheet(path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = '预算'
    book.active.append(['项目', '金额'])
    book.active.append(['GPU', 1200])
    book.active.append(['合计', '=B2'])
    book.create_sheet('空表')
    book.save(path)


@needs_textutil
@pytest.mark.parametrize('fmt', ['docx', 'doc', 'rtf', 'odt'])
def test_word_formats_are_read_through_textutil(tmp_path, fmt):
    word(tmp_path / f'计划.{fmt}', '第一章 研究计划\n用 BM25 检索论文。', fmt)
    out = read(library(tmp_path), f'计划.{fmt}')
    assert '用 BM25 检索论文' in out and 'evidence:' in out


def test_slides_are_units_with_tables_and_notes(tmp_path):
    slides(tmp_path / '汇报.pptx')
    sb = library(tmp_path)
    parts = extract.units(sb, tmp_path / '汇报.pptx')
    assert len(parts) == 2 and parts[0].startswith('【第 1 张幻灯片】')
    assert '规划、执行、核验三层' in parts[0] and '备注：讲到核验时停顿' in parts[0]
    assert 'Recall@5\t81.8%' in parts[1] and '分组里的说明' in parts[1]
    assert 'Agent 架构图' in read(sb, '汇报.pptx')


def test_sheets_are_units_of_values_not_formulas(tmp_path):
    sheet(tmp_path / '预算.xlsx')
    parts = extract.units(library(tmp_path), tmp_path / '预算.xlsx')
    assert parts[0] == '【工作表 预算】\n项目\t金额\nGPU\t1200\n合计' and parts[1] == '【工作表 空表】\n'


def test_long_sheets_are_cut_and_say_so(tmp_path, monkeypatch):
    from openpyxl import Workbook
    book = Workbook()
    for i in range(20): book.active.append([i])
    book.save(tmp_path / 'a.xlsx')
    monkeypatch.setattr(extract, 'MAX_ROWS', 5)
    assert extract.units(library(tmp_path), tmp_path / 'a.xlsx')[0].endswith('……（只抽取了前 5 行）')


def test_html_drops_tags_scripts_and_styles_and_rst_reads_as_text(tmp_path):
    (tmp_path / 'page.html').write_text('<html><head><title>标题</title><style>p{color:red}</style>'
                                        '<script>alert(1)</script></head><body><p>正文&amp;引用</p><div>第二段</div></body></html>')
    (tmp_path / 'guide.rst').write_text('Guide\n=====\n\nUse ``read_file``.')
    sb = library(tmp_path)
    out = read(sb, 'page.html')
    assert '正文&amp;引用' in out and '第二段' in out and 'alert' not in out and 'color' not in out
    assert '网页文字片段' in sb.artifacts.data['evidence'][out.split('evidence:')[1].split('；')[0]]['scope']
    assert 'Use ``read_file``' in read(sb, 'guide.rst')


def test_lock_damaged_and_oversized_files_get_clear_errors(tmp_path, monkeypatch):
    (tmp_path / '~$计划.docx').write_bytes(b'\x05owner')
    (tmp_path / 'broken.xlsx').write_bytes(b'not a zip')
    sheet(tmp_path / 'big.xlsx')
    sb = library(tmp_path)
    assert '锁文件' in read(sb, '~$计划.docx')
    assert '已加密或已损坏' in read(sb, 'broken.xlsx')
    monkeypatch.setattr(extract, 'MAX_UNZIPPED', 100)
    assert '解压后超过' in read(sb, 'big.xlsx')


def test_page_is_only_for_pdf(tmp_path):
    sheet(tmp_path / '预算.xlsx')
    assert 'page 只适用于 PDF' in read(library(tmp_path), '预算.xlsx', page=1)


def test_office_text_is_cached_by_content_and_refreshed_on_change(tmp_path, monkeypatch):
    sheet(tmp_path / '预算.xlsx')
    sb = library(tmp_path)
    first = extract.units(sb, tmp_path / '预算.xlsx')
    monkeypatch.setattr(extract, '_xlsx', lambda p: pytest.fail('cached text must not be re-extracted'))
    assert extract.units(Sandbox(tmp_path), tmp_path / '预算.xlsx') == first
    monkeypatch.undo()
    from openpyxl import Workbook
    book = Workbook(); book.active.append(['新内容']); book.save(tmp_path / '预算.xlsx')
    assert '新内容' in extract.units(Sandbox(tmp_path), tmp_path / '预算.xlsx')[0]


@needs_textutil
def test_when_sensitive_checks_are_on_a_docx_is_checked_by_its_text(tmp_path):
    # Before the shared extractor the rules read a docx's raw zip bytes and found nothing.
    from test_sensitive import id_number
    word(tmp_path / 'notes.docx', f'客户身份证 {id_number()}')
    assert '敏感文件' in read(library(tmp_path, sensitive_files=True), 'notes.docx')
    assert id_number() in read(library(tmp_path), 'notes.docx')


# User decision 2026-10-29: a PDF that needs a password is not read; say so plainly.
# Before: AES gave pypdf's "cryptography>=3.1 is required for AES algorithm", RC4
# "FileNotDecryptedError", neither saying the file has a password.

def pdf(path, **encryption):
    import pymupdf
    doc = pymupdf.open(); doc.new_page().insert_text((72, 72), 'Federated learning for hospital records.')
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path, **encryption)


@pytest.mark.parametrize('method', ['PDF_ENCRYPT_AES_256', 'PDF_ENCRYPT_RC4_128'])
def test_a_pdf_that_needs_a_password_is_not_read_and_says_why(tmp_path, method):
    import pymupdf
    pdf(tmp_path / '资料/locked.pdf', encryption=getattr(pymupdf, method), user_pw='u1', owner_pw='o1')
    sb = library(tmp_path)
    text = read(sb, '资料/locked.pdf')
    assert text.startswith('ERROR: locked.pdf 设有打开密码，FileButler 不读取加密 PDF 的内容')
    assert 'cryptography' not in text and 'Decrypted' not in text
    info = registry.dispatch('file_info', json.dumps({'path': '资料/locked.pdf'}), sb)
    assert '加密: 设有打开密码，不读取内容' in info and '读取失败' not in info


def test_an_owner_password_alone_does_not_stop_reading(tmp_path):
    import pymupdf
    pdf(tmp_path / '资料/owner.pdf', encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='o1', permissions=0)
    assert 'Federated learning' in read(library(tmp_path), '资料/owner.pdf')
