"""Text of a library file, as a list of units: one per PDF page, slide or sheet;
one for any other document. Every reader of contents (read_file, search_content,
paper_profiles, the sensitive-file rules) takes its text from here.

Formats (plan 2026-10-24, docs/feature-plan-2026-10-24.md):
- PDF: PyMuPDF, then pypdf (tools/fs_read._pdf_pages);
- plain text: UTF-8, then GBK;
- docx, doc, rtf, odt: macOS textutil (no extra package; 45 of 47 local Word files
  read, the 2 failures were a lock file and a text file named .docx);
- pptx: python-pptx, a unit per slide with its notes;
- xlsx: openpyxl (values, not formulas), a unit per sheet, rows as tab-separated lines;
- html: tags removed (stdlib html.parser);
- images, and PDF pages with almost no text layer: local OCR (agent/ocr.py).
Everything but plain text is cached by content (agent/document_cache).
"""
import importlib.metadata
import re
import shutil
import subprocess
import zipfile
from html.parser import HTMLParser

from agent import ocr
from tools.registry import ToolError

TEXT_EXTS = {".txt", ".md", ".csv", ".json", ".log", ".yaml", ".yml", ".py", ".rst"}
TEXTUTIL_EXTS = {".docx", ".doc", ".rtf", ".odt"}
OFFICE_EXTS = TEXTUTIL_EXTS | {".pptx", ".xlsx"}
HTML_EXTS = {".html", ".htm"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".heic", ".gif", ".bmp", ".webp"}
READABLE_EXTS = {".pdf"} | TEXT_EXTS | OFFICE_EXTS | HTML_EXTS | IMAGE_EXTS

VERSION = 1             # bump when an extractor's output changes
TEXTUTIL_TIMEOUT = 30   # seconds
MAX_UNZIPPED = 300 << 20  # an Office file larger than this unpacked is refused (zip bomb)
MAX_ROWS = 5000         # rows read per sheet
MAX_SHEET_CHARS = 200_000


def units(sb, path, *, start=1, pages=2):
    """Text units of path; PDFs honour start/pages, other formats are read whole."""
    from agent.document_cache import cached, pdf_pages
    from tools import fs_read
    ext = path.suffix.lower()
    if ext == ".pdf":
        found = pdf_pages(sb, path, start=start, pages=pages, extractor=fs_read._pdf_pages)
        return [_scanned(sb, path, start + i, text) if len(text.strip()) < ocr.MIN_TEXT else text
                for i, text in enumerate(found)]
    if ext in IMAGE_EXTS:
        if not ocr.budget(sb):
            ocr.stats(sb)["skipped"] += 1
            raise ToolError(f"{path.name} 本次没有时间识别文字，稍后重试")
        result = cached(sb, path, ["ocr-image", VERSION, ocr.version()], ocr.image_text)
        ocr.stats(sb)["images"] += 1
        return result
    if ext in TEXT_EXTS:
        return [decode(path.read_bytes())]
    if path.name.startswith("~$"):
        raise ToolError(f"{path.name} 是 Office 的锁文件（文档打开时自动生成），不是文档内容")
    if ext in HTML_EXTS:
        return cached(sb, path, ["html", VERSION], _html)
    if ext in OFFICE_EXTS:
        extractor = _textutil if ext in TEXTUTIL_EXTS else _pptx if ext == ".pptx" else _xlsx
        return cached(sb, path, ["office", ext, VERSION, _package_version(ext)], _guarded(extractor))
    raise ToolError(f"不支持读取 {ext or '无后缀'} 文件的内容，请用 file_info 查看基本信息")


def _scanned(sb, path, page, text):
    """A PDF page with (almost) no text layer: its OCR text, or the original text when
    OCR is unavailable, fails or is out of time (counted in ocr.stats)."""
    from agent.document_cache import cached
    counts = ocr.stats(sb)
    if not ocr.budget(sb):
        counts["skipped"] += 1
        return text
    try:
        found = cached(sb, path, ["pdf-ocr", page, VERSION, ocr.version()], lambda p: ocr.pdf_page_text(p, page))[0]
    except (ToolError, ValueError):
        counts["failed"] += 1
        return text
    counts["pages"] += 1
    return found if len(found.strip()) > len(text.strip()) else text


def decode(raw):
    for enc in ("utf-8", "gbk"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ToolError("无法识别编码（试过 UTF-8 和 GBK）")


def _package_version(ext):
    package = {".pptx": "python-pptx", ".xlsx": "openpyxl"}.get(ext)
    if not package:
        return "textutil"
    try:
        return importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:
        return "absent"


def _guarded(extractor):
    def run(path):
        if path.suffix.lower() in {".docx", ".pptx", ".xlsx", ".odt"}:
            if not zipfile.is_zipfile(path):
                raise ToolError(f"{path.name} 无法读取：文件已加密或已损坏")
            with zipfile.ZipFile(path) as z:
                if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                    raise ToolError(f"{path.name} 解压后超过 {MAX_UNZIPPED >> 20}MB，不读取")
        try:
            return extractor(path)
        except ToolError:
            raise
        except Exception as exc:
            raise ToolError(f"{path.name} 读取失败（{type(exc).__name__}），文件可能已加密或已损坏") from exc
    return run


def _textutil(path):
    tool = shutil.which("textutil")
    if not tool:
        raise ToolError(f"读取 {path.suffix} 需要 macOS 自带的 textutil，本机没有")
    try:
        done = subprocess.run([tool, "-convert", "txt", "-stdout", "--", str(path)],
                              capture_output=True, timeout=TEXTUTIL_TIMEOUT)
    except subprocess.TimeoutExpired:
        raise ToolError(f"{path.name} 读取超时（{TEXTUTIL_TIMEOUT} 秒）") from None
    if done.returncode:
        raise ToolError(f"{path.name} 无法读取：文件已加密、已损坏或不是 {path.suffix} 格式")
    return [done.stdout.decode("utf-8", errors="replace")]


def _pptx(path):
    from pptx import Presentation
    result = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        lines = [line for shape in slide.shapes for line in _shape_text(shape)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines += ["备注：" + notes.text]
        result.append(f"【第 {number} 张幻灯片】\n" + "\n".join(lines))
    return result


def _shape_text(shape):
    from pptx.shapes.group import GroupShape
    # Not shape.shape_type: it raises for shapes python-pptx does not recognise (seen in a real deck).
    if isinstance(shape, GroupShape):
        return [line for child in shape.shapes for line in _shape_text(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return ["\t".join(cell.text for cell in row.cells) for row in shape.table.rows]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
        return [shape.text_frame.text]
    return []


def _xlsx(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        result = []
        for sheet in book.worksheets:
            lines, size, cut = [], 0, False
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS or size > MAX_SHEET_CHARS:
                    cut = True
                    break
                cells = ["" if v is None else str(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    lines.append("\t".join(cells))
                    size += len(lines[-1]) + 1
            if cut:
                lines.append(f"……（只抽取了前 {len(lines)} 行）")
            result.append(f"【工作表 {sheet.title}】\n" + "\n".join(lines))
        return result
    finally:
        book.close()


class _Text(HTMLParser):
    SKIP = {"script", "style", "noscript", "template"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table", "pre"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skipping = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skipping:
            self.skipping -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skipping:
            self.parts.append(data)


def _html(path):
    parser = _Text()
    parser.feed(decode(path.read_bytes()))
    parser.close()
    text = re.sub(r"[ \t\r\f\v]+", " ", "".join(parser.parts))
    return [re.sub(r"\n\s*\n+", "\n\n", text).strip()]
